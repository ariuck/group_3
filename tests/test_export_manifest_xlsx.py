"""검수표 엑셀 내보내기(tools/export_manifest_xlsx.py) 시험.  openpyxl 이 없으면 건너뛴다.

실행 (프로젝트 폴더에서):  python -m unittest tests/test_export_manifest_xlsx.py
"""
import csv
import shutil
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
except ImportError:                                             # 선택 설치라 없을 수 있다
    openpyxl = None

from tools import export_manifest_xlsx as ex


@unittest.skipIf(openpyxl is None, "openpyxl 이 설치되어 있지 않음")
class ExportXlsxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.template = self.tmp / "틀.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "검수표"
        ws.append(["No", "이미지 파일명", "최종 BBox 수", "상태"])
        from openpyxl.worksheet.datavalidation import DataValidation
        dv = DataValidation(type="list", formula1='"검수 전,검수 완료"')
        dv.add("D2:D5")                                         # 틀은 5행까지만
        ws.add_data_validation(dv)
        wb.save(self.template)
        self.csv = self.tmp / "m.csv"
        with open(self.csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["No", "이미지 파일명", "최종 BBox 수", "상태"])
            for i in range(1, 9):
                w.writerow([i, f"p{i}.jpg", i, "검수 완료"])
        self.out = self.tmp / "out" / "m.xlsx"

    def test_rows_are_filled_with_numbers_as_numbers(self):
        n = ex.export(self.csv, self.template, self.out)
        self.assertEqual(n, 8)
        ws = openpyxl.load_workbook(self.out)["검수표"]
        self.assertEqual([c.value for c in ws[9]], [8, "p8.jpg", 8, "검수 완료"])
        self.assertIsInstance(ws["A2"].value, int)

    def test_choice_lists_are_extended_to_the_last_row(self):
        ex.export(self.csv, self.template, self.out)
        ws = openpyxl.load_workbook(self.out)["검수표"]
        self.assertEqual(str(ws.data_validations.dataValidation[0].sqref), "D2:D9")

    def test_template_is_not_modified(self):
        before = self.template.read_bytes()
        ex.export(self.csv, self.template, self.out)
        self.assertEqual(self.template.read_bytes(), before)

    def test_missing_column_is_reported(self):
        self.csv.write_text("No,다른칸\n1,x\n", encoding="utf-8-sig")
        with self.assertRaises(ValueError):
            ex.export(self.csv, self.template, self.out)


if __name__ == "__main__":
    unittest.main()
