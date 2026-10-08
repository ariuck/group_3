"""검수표(CSV)를 엑셀 틀(manifests/dataset_manifest.xlsx)에 채워서 제출용 xlsx 로 만든다.

    python tools/export_manifest_xlsx.py                      # data/final/dataset_manifest.xlsx 로 저장
    python tools/export_manifest_xlsx.py -o 내파일.xlsx

필요한 것: openpyxl  (pip install openpyxl — 이 도구에서만 쓰는 선택 설치라 requirements.txt 에는 넣지 않았다)
틀 파일(manifests/dataset_manifest.xlsx)은 고치지 않는다. 실제 사진 파일명이 들어가므로 결과 파일은 Git 에 올리지 않는다. (data/ 는 .gitignore)
"""
import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src import settings                                    # noqa: E402

TEMPLATE = settings.PROJECT_DIR / "manifests" / "dataset_manifest.xlsx"
INT_COLUMNS = ("No", "원본 BBox 수", "최종 BBox 수", "TXT 줄 수 = 화면 BBox 수")


def export(csv_path=None, template=None, out_path=None):
    """CSV 의 줄을 틀의 '검수표' 시트에 채워 out_path 로 저장하고 채운 줄 수를 돌려준다."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import MultiCellRange

    csv_path = Path(csv_path or settings.MANIFEST_PATH)
    template = Path(template or TEMPLATE)
    out_path = Path(out_path or settings.FINAL_DIR / "dataset_manifest.xlsx")
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    wb = openpyxl.load_workbook(template)
    ws = wb["검수표"]
    header = [c.value for c in ws[1]]
    missing = [h for h in header if h and rows and h not in rows[0]]
    if missing:
        raise ValueError(f"검수표 CSV 에 틀의 칸이 없습니다: {missing}")
    for i, r in enumerate(rows, 2):
        for j, h in enumerate(header, 1):
            v = r.get(h, "")
            if h in INT_COLUMNS and v.strip().lstrip("-").isdigit():
                v = int(v)
            ws.cell(row=i, column=j, value=v if v != "" else None)
    last = max(len(rows) + 1, ws.max_row)
    for dv in ws.data_validations.dataValidation:               # 선택 목록이 틀은 200행까지만 걸려 있어 마지막 줄까지 늘린다
        cols = sorted({rng.min_col for rng in dv.sqref.ranges})
        dv.sqref = MultiCellRange(" ".join(f"{get_column_letter(c)}2:{get_column_letter(c)}{last}" for c in cols))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return len(rows)


def main(argv=None):
    ap = argparse.ArgumentParser(description="검수표(CSV)를 엑셀 틀에 채워 xlsx 로 만든다")
    ap.add_argument("-o", "--output", default=None, help="저장할 파일 (기본: data/final/dataset_manifest.xlsx)")
    args = ap.parse_args(argv)
    try:
        import openpyxl                                         # noqa: F401
    except ImportError:
        print("openpyxl 이 필요합니다:  pip install openpyxl")
        return 2
    out = Path(args.output) if args.output else settings.FINAL_DIR / "dataset_manifest.xlsx"
    n = export(out_path=out)
    print(f"저장했습니다: {out}  (검수표 {n}줄)")
    if out.resolve().is_relative_to(settings.PROJECT_DIR / "data"):
        print("실제 사진 파일명이 들어 있어 Git 에 올리지 않는 위치(data/)입니다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
